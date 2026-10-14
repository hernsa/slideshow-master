"""Generic structural verifier for skill-built .pptx decks.

Usage: python verify-pptx.py <deck.pptx> [--min-rasters N] [--allow-text-only]

Fails on: wrong canvas, missing notes, raster floor breach, missing
transitions, auto-advance, click-animation trees beyond the entrance
budget (>3/slide or non-fade effects), zero morph continuity across
the whole deck, placeholder text.
"""
import math
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from lxml.etree import _Element


def xp(el, path):
    """lxml xpath with our namespaces (python-pptx overrides .xpath())."""
    return _Element.xpath(el, path, namespaces=NS)

NS = {
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "p14": "http://schemas.microsoft.com/office/powerpoint/2010/main",
    "p159": "http://schemas.microsoft.com/office/powerpoint/2015/09/main",
}

PLACEHOLDERS = ("lorem", "todo:", "tbd", "fixme", "slide title here")


def slide_rasters(slide):
    blobs = set()
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            try:
                blobs.add(shape.image.blob)
            except Exception:
                pass
    return blobs


def slide_text(slide):
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text)
    if slide.has_notes_slide:
        out.append(slide.notes_slide.placeholders[1].text)
    return "\n".join(out).lower()


def main():
    args = sys.argv[1:]
    flags = [a for a in args if a.startswith("--")]
    positional = [a for a in args if not a.startswith("--")]
    if not positional:
        print("usage: python verify-pptx.py <deck.pptx> "
              "[--min-rasters N] [--allow-text-only]")
        sys.exit(1)
    allow_text_only = "--allow-text-only" in flags
    min_rasters = None
    if "--min-rasters" in args:
        try:
            min_rasters = int(args[args.index("--min-rasters") + 1])
        except (ValueError, IndexError):
            print("FAIL: --min-rasters needs an integer")
            sys.exit(1)

    prs = Presentation(str(positional[0]))
    slides = list(prs.slides)
    total = len(slides)
    failures, warnings = [], []
    print("slides: %d" % total)
    if total == 0:
        print("FAIL: deck has no slides")
        sys.exit(1)

    # canvas
    from pptx.util import Inches
    if abs(prs.slide_width - Inches(13.333)) > Inches(0.01) or \
       abs(prs.slide_height - Inches(7.5)) > Inches(0.01):
        failures.append("canvas is not 13.333x7.5in "
                        "(got %.3fx%.3fin)"
                        % (prs.slide_width / 914400,
                           prs.slide_height / 914400))

    all_blobs = set()
    raster_slides = 0
    prev_names = None
    continuity_pairs = 0
    for i, slide in enumerate(slides, start=1):
        el = slide._element
        blobs = slide_rasters(slide)
        all_blobs |= blobs
        if blobs:
            raster_slides += 1
        # notes
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.placeholders[1].text.strip()
        if not notes:
            failures.append("slide %d: no speaker notes" % i)
        # transitions (p:transition, p14/p159 morph, or fade fallback)
        has_trans = bool(xp(el, ".//p:transition|.//p:fade|"
                              ".//p14:morph|.//p159:morph"))
        if not has_trans:
            failures.append("slide %d: no transition authored" % i)
        # manual advance: advTm must be absent
        for t in xp(el, ".//p:transition"):
            if t.get("advTm") is not None:
                failures.append("slide %d: auto-advance set "
                                "(advTm present)" % i)
        # timing tree: bounded entrance budget only
        fx = xp(el, ".//p:timing//p:animEffect")
        if len(fx) > 3:
            failures.append("slide %d: %d entrance effects "
                            "(budget is 3)" % (i, len(fx)))
        for f in fx:
            if (f.get("filter") or "").lower() != "fade":
                failures.append("slide %d: non-fade entrance "
                                "filter=%r" % (i, f.get("filter")))
        # continuity
        names = {s.name for s in slide.shapes if s.name}
        if prev_names is not None:
            if names & prev_names:
                continuity_pairs += 1
            else:
                warnings.append("slides %d-%d share no shape names "
                                "(morph has nothing to carry)" % (i - 1, i))
        prev_names = names
        # placeholders
        txt = slide_text(slide)
        for ph in PLACEHOLDERS:
            if ph in txt:
                failures.append("slide %d: placeholder text %r" % (i, ph))
        print("slide %d: rasters=%d notes=%s trans=%s fx=%d" %
              (i, len(blobs), "yes" if notes else "NO",
               "yes" if has_trans else "NO", len(fx)))

    # raster floor
    raster_total = len(all_blobs)
    target = min_rasters if min_rasters is not None else total
    print("Rasters: %d unique across %d/%d slides (target >= %d)%s"
          % (raster_total, raster_slides, total, target,
             " (text-only deck, raster gate off)"
             if allow_text_only else ""))
    if not allow_text_only:
        if raster_total == 0:
            failures.append("zero rasters in the deck — %d slides reference "
                            "no photographic image. Every slide needs at "
                            "least one raster (pptx-motion-language.md "
                            "section 5); pass --allow-text-only only for a "
                            "genuinely type-only deck." % total)
        elif raster_total < target:
            failures.append("raster count %d below target %d "
                            "(visual-plan.md hard counts)" % (raster_total,
                                                              target))
        need = min(total, math.ceil(total * 0.6))
        if raster_slides < need:
            failures.append("rasters on %d/%d slides (need %d): a slide "
                            "with no image reads as a slide sorter"
                            % (raster_slides, total, need))
    if continuity_pairs == 0 and total > 1:
        failures.append("no shared shape names across any consecutive pair "
                        "— the deck is a slide sorter, Morph carries nothing")
    for w in warnings:
        print("warn:", w)
    if failures:
        print("FAIL")
        for f in failures:
            print(" -", f)
        sys.exit(1)
    print("OK: deck structure passes")


if __name__ == "__main__":
    main()
